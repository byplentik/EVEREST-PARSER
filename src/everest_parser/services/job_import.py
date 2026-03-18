"""Импорт входного `.xlsx` в batch-job и очередь задач."""

from __future__ import annotations

from dataclasses import asdict
from dataclasses import dataclass
from datetime import datetime
from datetime import timezone
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from everest_parser.db import JobStatus
from everest_parser.db import ParseJob
from everest_parser.db import ParseTask
from everest_parser.db import ParserType
from everest_parser.db import TaskStatus
from everest_parser.db import session_scope


HEADER_ALIASES: dict[ParserType, tuple[str, ...]] = {
    ParserType.FEDRESURS: ("inn", "инн"),
    ParserType.KAD: ("case_number", "номер_дела", "дело"),
}

HEADER_SEPARATOR_RE = re.compile(r"[\s\-]+")
MULTISPACE_RE = re.compile(r"\s+")
DIGITS_ONLY_RE = re.compile(r"^\d+$")
DASH_TRANSLATION_TABLE = str.maketrans(
    {
        "–": "-",
        "—": "-",
        "−": "-",
        "―": "-",
        "／": "/",
        "\\": "/",
    }
)


class JobImportError(Exception):
    """Ошибка входного pipeline при чтении файла или подготовке задач."""


@dataclass(slots=True)
class PreparedTask:
    """Подготовленная к записи задача."""

    source_value: str
    normalized_value: str


@dataclass(slots=True)
class ImportSummary:
    """Сводка по результату импорта `.xlsx`."""

    job_id: int
    job_status: str
    parser_type: str
    source_file: str
    worksheet_title: str
    source_column: str
    rows_read: int
    empty_rows: int
    invalid_rows: int
    duplicate_rows: int
    tasks_created: int

    def as_payload(self) -> dict[str, object]:
        """Вернуть JSON-совместимый словарь для CLI-ответа."""

        return asdict(self)


@dataclass(slots=True)
class PreparedImport:
    """Результат чтения и нормализации входного файла."""

    source_file: str
    worksheet_title: str
    source_column: str
    rows_read: int
    empty_rows: int
    invalid_rows: int
    duplicate_rows: int
    tasks: list[PreparedTask]


def create_job_from_xlsx(parser_type: ParserType, input_path: str | Path) -> ImportSummary:
    """Создать job и задачи в БД на основе входного `.xlsx` файла."""

    file_path = validate_input_path(input_path)
    prepared_import = read_tasks_from_xlsx(parser_type=parser_type, input_path=file_path)
    return persist_import(parser_type=parser_type, prepared_import=prepared_import)


def validate_input_path(input_path: str | Path) -> Path:
    """Проверить, что входной путь существует и указывает на `.xlsx` файл."""

    file_path = Path(input_path).expanduser()

    if not file_path.exists():
        raise JobImportError(f"Файл не найден: {file_path}")

    if not file_path.is_file():
        raise JobImportError(f"Путь не указывает на файл: {file_path}")

    if file_path.suffix.lower() != ".xlsx":
        raise JobImportError("Поддерживаются только файлы формата .xlsx")

    return file_path


def read_tasks_from_xlsx(parser_type: ParserType, input_path: Path) -> PreparedImport:
    """Прочитать `.xlsx`, нормализовать значения и подготовить список задач."""

    try:
        workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
    except FileNotFoundError as error:
        raise JobImportError(f"Файл не найден: {input_path}") from error
    except (BadZipFile, InvalidFileException, OSError) as error:
        raise JobImportError(f"Не удалось прочитать .xlsx файл: {input_path}") from error

    try:
        if not workbook.sheetnames:
            raise JobImportError("Входной .xlsx не содержит листов")

        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)

        if header_row is None:
            raise JobImportError("Входной .xlsx не содержит строк")

        column_index, source_column = resolve_source_column(parser_type, header_row)

        prepared_tasks: list[PreparedTask] = []
        seen_normalized_values: set[str] = set()
        rows_read = 0
        empty_rows = 0
        invalid_rows = 0
        duplicate_rows = 0

        for row in rows:
            rows_read += 1
            raw_value = row[column_index] if column_index < len(row) else None
            source_value = stringify_cell(raw_value)

            if not source_value:
                empty_rows += 1
                continue

            normalized_value = normalize_identifier(parser_type, source_value)
            if normalized_value is None:
                invalid_rows += 1
                continue

            if normalized_value in seen_normalized_values:
                duplicate_rows += 1
                continue

            seen_normalized_values.add(normalized_value)
            prepared_tasks.append(
                PreparedTask(
                    source_value=source_value,
                    normalized_value=normalized_value,
                )
            )

        return PreparedImport(
            source_file=input_path.as_posix(),
            worksheet_title=worksheet.title,
            source_column=source_column,
            rows_read=rows_read,
            empty_rows=empty_rows,
            invalid_rows=invalid_rows,
            duplicate_rows=duplicate_rows,
            tasks=prepared_tasks,
        )
    finally:
        workbook.close()


def resolve_source_column(parser_type: ParserType, header_row: tuple[object, ...]) -> tuple[int, str]:
    """Найти индекс колонки с идентификатором для выбранного парсера."""

    normalized_headers = [normalize_header(value) for value in header_row]
    aliases = HEADER_ALIASES[parser_type]

    for alias in aliases:
        if alias in normalized_headers:
            return normalized_headers.index(alias), alias

    available_headers = [header for header in normalized_headers if header]
    raise JobImportError(
        "Входной .xlsx не содержит обязательной колонки. "
        f"Ожидались: {', '.join(aliases)}. "
        f"Найдено: {', '.join(available_headers) or 'нет заголовков'}."
    )


def normalize_header(value: object) -> str:
    """Нормализовать имя колонки для поиска по алиасам."""

    text = stringify_cell(value).lower()
    text = HEADER_SEPARATOR_RE.sub("_", text)
    return text.strip("_")


def stringify_cell(value: object) -> str:
    """Преобразовать значение ячейки в строку без лишних пробелов."""

    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value).strip()

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")

    return str(value).strip()


def normalize_identifier(parser_type: ParserType, source_value: str) -> str | None:
    """Нормализовать идентификатор в зависимости от типа парсера."""

    if parser_type is ParserType.FEDRESURS:
        return normalize_inn(source_value)

    if parser_type is ParserType.KAD:
        return normalize_case_number(source_value)

    raise JobImportError(f"Неподдерживаемый тип парсера: {parser_type}")


def normalize_inn(source_value: str) -> str | None:
    """Нормализовать ИНН и проверить его базовый формат."""

    normalized = source_value.replace("\u00a0", "")
    normalized = MULTISPACE_RE.sub("", normalized)

    if not normalized:
        return None

    if not DIGITS_ONLY_RE.fullmatch(normalized):
        return None

    if len(normalized) not in (10, 12):
        return None

    return normalized


def normalize_case_number(source_value: str) -> str | None:
    """Нормализовать номер дела для `kad.arbitr.ru`."""

    normalized = source_value.replace("\u00a0", " ")
    normalized = normalized.translate(DASH_TRANSLATION_TABLE)
    normalized = normalized.upper().strip()
    normalized = MULTISPACE_RE.sub("", normalized)

    if not normalized:
        return None

    if "/" not in normalized:
        return None

    return normalized


def determine_job_status(prepared_import: PreparedImport) -> tuple[JobStatus, datetime | None]:
    """Определить стартовый статус job после импорта входного файла."""

    if prepared_import.tasks:
        return JobStatus.PENDING, None

    if prepared_import.invalid_rows or prepared_import.duplicate_rows:
        return JobStatus.COMPLETED_WITH_ERRORS, datetime.now(timezone.utc)

    return JobStatus.COMPLETED, datetime.now(timezone.utc)


def persist_import(parser_type: ParserType, prepared_import: PreparedImport) -> ImportSummary:
    """Сохранить job и задачи в PostgreSQL."""

    job_status, finished_at = determine_job_status(prepared_import)

    with session_scope() as session:
        job = ParseJob(
            parser_type=parser_type,
            source_file=prepared_import.source_file,
            status=job_status,
            total_items=len(prepared_import.tasks),
            finished_at=finished_at,
        )
        session.add(job)
        session.flush()

        if prepared_import.tasks:
            session.add_all(
                [
                    ParseTask(
                        job_id=job.id,
                        source_value=task.source_value,
                        normalized_value=task.normalized_value,
                        status=TaskStatus.PENDING,
                    )
                    for task in prepared_import.tasks
                ]
            )

        session.flush()

        return ImportSummary(
            job_id=job.id,
            job_status=job.status.value,
            parser_type=parser_type.value,
            source_file=prepared_import.source_file,
            worksheet_title=prepared_import.worksheet_title,
            source_column=prepared_import.source_column,
            rows_read=prepared_import.rows_read,
            empty_rows=prepared_import.empty_rows,
            invalid_rows=prepared_import.invalid_rows,
            duplicate_rows=prepared_import.duplicate_rows,
            tasks_created=len(prepared_import.tasks),
        )
